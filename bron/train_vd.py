"""Harmoniseer de 5 jaartabbladen van het Randwijk-weerstation tot 1 nette tijdreeks."""
import pandas as pd, numpy as np
from datetime import datetime, time, timedelta
from openpyxl import load_workbook

import os, sys
SRC = sys.argv[1] if len(sys.argv) > 1 else "Weerdata_FRCRandwijk_2021_tm_2025.xlsx"

# Kolom-index -> standaardnaam, per tabblad (0-indexed)
MAPS = {
    "2021": {2:"temp",3:"temp_hi",5:"hum",6:"dewpt",7:"wind",8:"winddir",16:"rain",17:"rainrate",33:"leafwet1",34:"leafwet2"},
    "2022": {3:"temp",4:"temp_hi",5:"hum",6:"wind",7:"winddir",8:"rain",9:"rainrate",18:"leafwet1",19:"leafwet2"},
    "2023": {4:"temp",5:"temp_hi",6:"temp_lo",7:"hum",8:"dewpt",10:"wind",11:"winddir",18:"rain",19:"rainrate",31:"leafwet1",32:"leafwet2"},
    "2024": {3:"temp",4:"temp_hi",5:"hum",6:"wind",7:"winddir",8:"rain",9:"rainrate",10:"leafwet1",11:"leafwet2"},
    "2025": {3:"temp",4:"temp_hi",5:"hum",6:"wind",7:"winddir",8:"rain",9:"rainrate",18:"leafwet1",19:"leafwet2"},
}
# Tijdkolommen: (tijd_index, ampm_index of None voor 24u)
TIMECOL = {"2021":(1,None),"2022":(1,2),"2023":(1,2),"2024":(1,2),"2025":(1,2)}

def to_num(v):
    if v is None: return np.nan
    if isinstance(v,(int,float)): return float(v)
    s = str(v).strip().replace(",",".")
    if s in ("","---","--","n/a","NA"): return np.nan
    try: return float(s)
    except: return np.nan

def build_dt(date_v, t_v, ampm):
    if date_v is None or t_v is None: return pd.NaT
    d = date_v.date() if isinstance(date_v, datetime) else date_v
    if isinstance(t_v, time):
        h,m = t_v.hour, t_v.minute
    else:
        f = float(t_v); h = int(f*24); m = int(round((f*24-h)*60))
    if ampm is not None:
        a = str(ampm).strip().upper()
        if a == "AM":
            if h == 12: h = 0
        elif a == "PM":
            if h != 12: h += 12
    return datetime(d.year, d.month, d.day, h, m)

frames = []
wb = load_workbook(SRC, read_only=True, data_only=True)
for sheet, colmap in MAPS.items():
    ws = wb[sheet]
    ti, ai = TIMECOL[sheet]
    recs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None: continue
        dt = build_dt(row[0], row[ti], row[ai] if ai is not None else None)
        if pd.isna(dt): continue
        rec = {"dt": dt}
        for idx, name in colmap.items():
            rec[name] = to_num(row[idx]) if idx < len(row) else np.nan
        recs.append(rec)
    df = pd.DataFrame(recs)
    df["bron_jaar"] = sheet
    frames.append(df)
    print(f"{sheet}: {len(df):>6} rijen  | temp NaN%={df['temp'].isna().mean()*100:4.1f}  hum NaN%={df['hum'].isna().mean()*100:4.1f}")

data = pd.concat(frames, ignore_index=True)
data = data.dropna(subset=["dt"]).sort_values("dt").drop_duplicates(subset=["dt"]).reset_index(drop=True)

# Plausibiliteitsfilters (sensorfouten eruit)
data.loc[(data.temp < -30) | (data.temp > 45), "temp"] = np.nan
data.loc[(data.hum < 1) | (data.hum > 100), "hum"] = np.nan
data.loc[data.wind < 0, "wind"] = np.nan

# Dauwpunt (Magnus) waar niet aanwezig, en VPD
def dewpoint(T, RH):
    a,b = 17.62, 243.12
    g = np.log(np.clip(RH,1,100)/100.0) + a*T/(b+T)
    return b*g/(a-g)
def svp(T):  # verzadigingsdampdruk kPa
    return 0.6108*np.exp(17.27*T/(T+237.3))
data["dewpt_calc"] = dewpoint(data.temp, data.hum)
data["dewpt"] = data["dewpt"].fillna(data["dewpt_calc"]) if "dewpt" in data else data["dewpt_calc"]
data["vpd"] = (svp(data.temp)*(1-np.clip(data.hum,0,100)/100.0))  # kPa

print("\n=== Overzicht samengevoegd ===")
print("Totaal metingen:", len(data))
print("Periode:", data.dt.min(), "->", data.dt.max())
print("Kolommen:", [c for c in data.columns if c not in ('dt','bron_jaar','dewpt_calc')])
print("\nTemp beschrijving:\n", data["temp"].describe().round(2).to_string())
print("\nMetingen onder 0C:", int((data.temp<0).sum()), f"({(data.temp<0).mean()*100:.1f}%)")

data.to_parquet("../data/weer_clean.parquet")
print("\nOpgeslagen -> weer_clean.parquet")
